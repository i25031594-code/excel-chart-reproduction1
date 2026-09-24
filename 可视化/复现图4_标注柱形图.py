"""复现图4：多色标注柱形图。兼容 Jupyter 和普通 .py 文件。"""
from pathlib import Path
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from openpyxl import load_workbook
EXCEL_FILE="第二章 图表(前15).xlsx"; SHEET="4 标注柱形图"; OUTPUT="图4_标注柱形图.png"; BG="#1A1E43"; WHITE="#FFFFFF"
COLORS=["#7BBDD5","#49A098","#E66B4C","#FFC000","#0097E0","#0070C0","#4A5BD1","#464CAC"]
def base_dir(): return Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
def excel_path():
    for p in (base_dir()/EXCEL_FILE,Path.cwd()/EXCEL_FILE,Path.cwd()/"upload"/EXCEL_FILE):
        if p.exists(): return p
    raise FileNotFoundError(f"未找到{EXCEL_FILE}")
def font():
    for p in ["C:/Windows/Fonts/msyh.ttc","C:/Windows/Fonts/simhei.ttf","/System/Library/Fonts/PingFang.ttc","/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]:
        if Path(p).exists(): return FontProperties(fname=p)
    return FontProperties(family=["Microsoft YaHei","SimHei","sans-serif"])
def read_data():
    wb=load_workbook(excel_path(),data_only=True,read_only=True); ws=wb[SHEET]; n=[ws.cell(r,2).value for r in range(3,11)]; v=np.array([ws.cell(r,3).value for r in range(3,11)],float); wb.close(); return n,v
def main():
    n,v=read_data(); fp=font(); x=np.arange(8); fig=plt.figure(figsize=(11.56,8.08),dpi=150,facecolor=BG); ax=fig.add_axes([.08,.16,.84,.55],facecolor=BG)
    bars=ax.bar(x,v,.50,color=COLORS,edgecolor="none",zorder=3); ax.set_ylim(0,10000); ax.set_xticks(x,n); ax.set_yticks(np.arange(0,10001,2000)); ax.tick_params(length=0,colors=WHITE,pad=7)
    for t in list(ax.get_xticklabels())+list(ax.get_yticklabels()): t.set_fontproperties(fp); t.set_fontsize(9)
    ax.set_axisbelow(True); ax.grid(axis="y",color="white",alpha=.17,linestyle=(0,(6,4)),lw=.8)
    for s in ax.spines.values(): s.set_visible(False)
    for b,y,c in zip(bars,v,COLORS): ax.text(b.get_x()+b.get_width()/2,y+150,f"{int(y)}",ha="center",color=WHITE,fontsize=8,fontproperties=fp,bbox=dict(boxstyle="square,pad=.12",fc=c,ec="none"))
    fig.text(.06,.92,"2021年商品销量情况",color=WHITE,fontsize=20,fontweight="bold",fontproperties=fp,va="top")
    fig.text(.06,.82,"口红销量最好达9221，是眼影最低值2645近3.5倍",color=WHITE,fontsize=14,fontproperties=fp,va="top")
    fig.text(.045,.045,"*注：数据来源于公司销售系统，统计日期截至2022.08.31",color="#D9D9D9",fontsize=8,fontproperties=fp)
    out=base_dir()/OUTPUT; fig.savefig(out,dpi=300,facecolor=BG); print(f"图片已保存：{out.resolve()}"); plt.show()
if __name__=="__main__": main()
