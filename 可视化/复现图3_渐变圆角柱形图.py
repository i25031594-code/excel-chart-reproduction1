"""复现图3：渐变圆角柱形图。兼容 Jupyter 和普通 .py 文件。"""
from pathlib import Path
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.font_manager import FontProperties
from matplotlib.patches import FancyBboxPatch
from openpyxl import load_workbook

EXCEL_FILE="第二章 图表(前15).xlsx"; SHEET="3 渐变圆角柱形图"; OUTPUT="图3_渐变圆角柱形图.png"
BG,WHITE="#1A1E43","#FFFFFF"
def base_dir(): return Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
def excel_path():
    for p in (base_dir()/EXCEL_FILE,Path.cwd()/EXCEL_FILE,Path.cwd()/"upload"/EXCEL_FILE):
        if p.exists(): return p
    raise FileNotFoundError(f"未找到{EXCEL_FILE}")
def font():
    for p in ["C:/Windows/Fonts/msyh.ttc","C:/Windows/Fonts/simhei.ttf","/System/Library/Fonts/PingFang.ttc","/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]:
        if Path(p).exists(): return FontProperties(fname=p)
    return FontProperties(family=["Microsoft YaHei","SimHei","sans-serif"])
def read_data():
    wb=load_workbook(excel_path(),data_only=True,read_only=True); ws=wb[SHEET]
    a=[ws.cell(r,2).value for r in range(3,9)]; b=np.array([ws.cell(r,3).value for r in range(3,9)],float); wb.close(); return a,b
def rounded_gradient(ax,x,v,width=.16):
    cmap=LinearSegmentedColormap.from_list("cyan",["#0070C0","#10E3E3"]); grad=np.linspace(0,1,400)[:,None]; patches=[]
    for xi,yi in zip(x,v):
        p=FancyBboxPatch((xi-width/2,0),width,yi,boxstyle=f"round,pad=0,rounding_size={width/2}",fc="none",ec="none",transform=ax.transData,zorder=3); ax.add_patch(p)
        im=ax.imshow(grad,extent=[xi-width/2,xi+width/2,0,yi],origin="lower",aspect="auto",cmap=cmap,zorder=3); im.set_clip_path(p); patches.append(p)
    return patches
def main():
    names,y=read_data(); x=np.arange(6); fp=font(); fig=plt.figure(figsize=(11.56,8.10),dpi=150,facecolor=BG); ax=fig.add_axes([.08,.15,.84,.57],facecolor=BG)
    rounded_gradient(ax,x,y); ax.set_xlim(-.5,5.5); ax.set_ylim(0,1200); ax.set_xticks(x,names); ax.set_yticks(np.arange(0,1201,200)); ax.tick_params(length=0,colors=WHITE,pad=7)
    for t in list(ax.get_xticklabels())+list(ax.get_yticklabels()): t.set_fontproperties(fp); t.set_fontsize(9)
    ax.set_axisbelow(True); ax.grid(axis="y",color="white",alpha=.16,linestyle=(0,(6,4)),lw=.8)
    for s in ax.spines.values(): s.set_visible(False)
    for xi,yi in zip(x,y): ax.text(xi,yi+70,f"{int(yi)}",ha="center",color="#9BCBFF",fontsize=9,fontproperties=fp,bbox=dict(boxstyle="round,pad=.55,rounding_size=.2",fc="#123A6E",ec="none",alpha=.96))
    fig.text(.06,.92,"3月商品销量对比",color=WHITE,fontsize=20,fontweight="bold",fontproperties=fp,va="top")
    fig.text(.06,.82,"防晒销量最多，3月销量856；面膜最少，3月销量523",color=WHITE,fontsize=14,fontproperties=fp,va="top")
    fig.text(.045,.045,"*注：数据来源于公司销售系统，统计日期截至2022.03.31",color="#D9D9D9",fontsize=8,fontproperties=fp)
    out=base_dir()/OUTPUT; fig.savefig(out,dpi=300,facecolor=BG); print(f"图片已保存：{out.resolve()}"); plt.show()
if __name__=="__main__": main()
