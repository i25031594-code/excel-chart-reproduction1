"""
复现《第二章 图表(前15).xlsx》中的第 1 个图：3月各区域销量分布。

使用方法：
1. 将本脚本与 Excel 文件放在同一文件夹；
2. 安装依赖：pip install openpyxl matplotlib numpy
3. 运行脚本：python 复现图1_渐变柱形图.py

输出：图1_3月各区域销量分布.png（300 dpi）
"""

from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.font_manager import FontProperties
from openpyxl import load_workbook


# ======================== 可修改参数 ========================
EXCEL_FILE = "第二章 图表(前15).xlsx"
SHEET_NAME = "1 渐变柱形图"
OUTPUT_FILE = "图1_3月各区域销量分布.png"

# 原 Excel 图的颜色与文字
BACKGROUND = "#1A1E43"
GRADIENT_TOP = "#0070C0"
GRADIENT_BOTTOM = "#00B0F0"
TEXT_COLOR = "#FFFFFF"
FOOTNOTE_COLOR = "#D9D9D9"
GRID_COLOR = "#FFFFFF"


def get_working_directory() -> Path:
    """兼容 Jupyter Notebook 和普通 .py 文件的所在目录。"""
    # Jupyter 中没有 __file__，此时 Path.cwd() 就是 Notebook 当前工作目录。
    if "__file__" in globals():
        return Path(__file__).resolve().parent
    return Path.cwd()


def find_excel_file(filename: str) -> Path:
    """依次在 Notebook/脚本目录、当前目录和 upload 子目录寻找 Excel。"""
    working_dir = get_working_directory()
    candidates = [
        working_dir / filename,
        Path.cwd() / filename,
        Path.cwd() / "upload" / filename,
    ]
    for path in candidates:
        if path.exists():
            return path
    raise FileNotFoundError(
        f"未找到 {filename}。请把本脚本与 Excel 文件放在同一文件夹。"
    )


def get_chinese_font() -> FontProperties:
    """优先使用原图的微软雅黑；没有时自动选择常见中文字体。"""
    font_paths = [
        # Windows
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/msyhbd.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        # macOS
        Path("/System/Library/Fonts/PingFang.ttc"),
        # Linux / 常见 Jupyter 环境
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJKSC-Regular.otf"),
        Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
    ]
    for font_path in font_paths:
        if font_path.exists():
            return FontProperties(fname=str(font_path))

    # 如果系统中没有上述字体，仍可运行；Windows 通常能直接找到微软雅黑。
    return FontProperties(
        family=["Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "sans-serif"]
    )


def read_chart_data(excel_path: Path):
    """严格读取第一个工作表 B3:C8 的区域与销售量。"""
    workbook = load_workbook(excel_path, data_only=True, read_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        raise KeyError(f"Excel 中不存在工作表：{SHEET_NAME}")

    sheet = workbook[SHEET_NAME]
    regions = [sheet.cell(row=row, column=2).value for row in range(3, 9)]
    sales = [sheet.cell(row=row, column=3).value for row in range(3, 9)]
    workbook.close()

    if any(value is None for value in regions + sales):
        raise ValueError("B3:C8 中存在空值，无法完整复现第一个图。")

    return [str(value) for value in regions], np.asarray(sales, dtype=float)


def add_gradient_bars(ax, x, heights, width=0.313):
    """绘制与原图一致的竖直渐变柱：上深蓝、下亮蓝。"""
    bars = ax.bar(x, heights, width=width, color="none", edgecolor="none", zorder=3)
    cmap = LinearSegmentedColormap.from_list(
        "excel_blue_gradient", [GRADIENT_BOTTOM, GRADIENT_TOP]
    )
    gradient = np.linspace(0, 1, 512).reshape(512, 1)

    for bar in bars:
        left = bar.get_x()
        right = left + bar.get_width()
        top = bar.get_height()
        image = ax.imshow(
            gradient,
            extent=[left, right, 0, top],
            origin="lower",
            aspect="auto",
            cmap=cmap,
            interpolation="bicubic",
            zorder=3,
        )
        image.set_clip_path(bar)

    return bars


def create_chart(regions, sales, output_path: Path):
    font = get_chinese_font()

    # 原图宽高约为 578:397；深色背景覆盖整张图片。
    fig = plt.figure(figsize=(11.56, 7.94), dpi=150, facecolor=BACKGROUND)

    # 对应 Excel 图表的手动绘图区：x≈14.4%、y≈31.9%、w≈76.0%、h≈52.5%。
    ax = fig.add_axes([0.144, 0.156, 0.760, 0.525], facecolor=BACKGROUND)

    x = np.arange(len(regions))
    bars = add_gradient_bars(ax, x, sales)

    # 坐标范围、刻度与虚线网格均按源图设置。
    ax.set_xlim(-0.5, len(regions) - 0.5)
    ax.set_ylim(0, 4000)
    ax.set_xticks(x)
    ax.set_xticklabels(regions, color=TEXT_COLOR, fontsize=9)
    ax.set_yticks([0, 1000, 2000, 3000, 4000])
    ax.set_yticklabels(["0", "1000", "2000", "3000", "4000"],
                       color=TEXT_COLOR, fontsize=9)

    for label in list(ax.get_xticklabels()) + list(ax.get_yticklabels()):
        label.set_fontproperties(font)
        label.set_fontsize(9)

    ax.tick_params(axis="both", length=0, colors=TEXT_COLOR, pad=6)
    ax.set_axisbelow(True)
    ax.yaxis.grid(
        True,
        color=GRID_COLOR,
        alpha=0.20,
        linewidth=0.8,
        linestyle=(0, (8, 4)),
    )
    ax.xaxis.grid(False)

    # 原图没有坐标轴边框。
    for spine in ax.spines.values():
        spine.set_visible(False)

    # 柱顶销售量标签。
    for bar, value in zip(bars, sales):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            value + 80,
            f"{int(value)}",
            ha="center",
            va="bottom",
            color=TEXT_COLOR,
            fontsize=9,
            fontproperties=font,
            zorder=5,
        )

    # 标题、副标题和脚注均为原图原文。
    fig.text(
        0.060,
        0.915,
        "3月各区域销量分布",
        ha="left",
        va="top",
        color=TEXT_COLOR,
        fontsize=20,
        fontweight="bold",
        fontproperties=font,
    )
    fig.text(
        0.060,
        0.815,
        "东北销量最多占比总销量的22%，华南销量最低",
        ha="left",
        va="top",
        color=TEXT_COLOR,
        fontsize=14,
        fontproperties=font,
    )
    fig.text(
        0.042,
        0.045,
        "*注：数据来源于公司销售系统，统计日期截至2022.03.31",
        ha="left",
        va="bottom",
        color=FOOTNOTE_COLOR,
        fontsize=8,
        fontproperties=font,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(
        output_path,
        dpi=300,
        facecolor=fig.get_facecolor(),
        edgecolor="none",
        bbox_inches=None,
        pad_inches=0,
    )
    print(f"图片已保存：{output_path.resolve()}")
    plt.show()


if __name__ == "__main__":
    excel_path = find_excel_file(EXCEL_FILE)
    chart_regions, chart_sales = read_chart_data(excel_path)
    output_file = get_working_directory() / OUTPUT_FILE
    create_chart(chart_regions, chart_sales, output_file)
