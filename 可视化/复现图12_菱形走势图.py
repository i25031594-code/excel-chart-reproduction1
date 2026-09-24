"""复现图12：公司计划完成率菱形走势图。"""
from pathlib import Path
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from openpyxl import load_workbook
EXCEL_FILE="第二章 图表(前15).xlsx"; SHEET="12 菱形走势图"; OUTPUT="图12_菱形走势图.png"; BG="#1A1E43"; WHITE="#FFFFFF"; RED="#E74E69"
def base_dir(): return Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
def excel_path():
    for p in (base_dir()/EXCEL_FILE,Path.cwd()/EXCEL_FILE,Path.cwd()/"upload"/EXCEL_FILE):
        if p.exists(): return p
    raise FileNotFoundError(f"未找到{EXCEL_FILE}")
def font():
    for p in ["C:/Windows/Fonts/msyh.ttc","C:/Windows/Fonts/simhei.ttf","/System/Library/Fonts/PingFang.ttc","/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]:
        if Path(p).exists(): return FontProperties(fname=p)
    return FontProperties(family=["Microsoft YaHei","SimHei","sans-serif"])
def read_data():
    wb=load_workbook(excel_path(),data_only=True,read_only=True); ws=wb[SHEET]; m=[ws.cell(r,2).value for r in range(3,11)]; v=np.array([ws.cell(r,3).value for r in range(3,11)],float)*100; wb.close(); return m,v
def main():
    months,v=read_data(); fp=font(); x=np.arange(8); fig=plt.figure(figsize=(11.56,7.58),dpi=150,facecolor=BG); ax=fig.add_axes([.08,.15,.86,.55],facecolor=BG)
    ax.vlines(x,0,v,color=RED,lw=.9,alpha=.9); ax.scatter(x,v,marker="D",s=70,facecolor=BG,edgecolor=RED,lw=1.4,zorder=3)
    ax.set_xlim(-.5,7.5); ax.set_ylim(0,82); ax.set_xticks(x,months); ax.set_yticks([]); ax.tick_params(axis="x",length=0,colors=WHITE,pad=8)
    for t in ax.get_xticklabels(): t.set_fontproperties(fp); t.set_fontsize(9)
    for s in ax.spines.values(): s.set_visible(False)
    ax.spines["bottom"].set_visible(True); ax.spines["bottom"].set_color("#A7ABC0"); ax.spines["bottom"].set_alpha(.55)
    for xi,yi in zip(x,v): ax.text(xi,yi+4,f"{yi:.2f}%",ha="center",color=WHITE,fontsize=9,fontproperties=fp)
    fig.text(.06,.92,"2022年1-8月公司计划完成率",color=WHITE,fontsize=20,fontweight="bold",fontproperties=fp,va="top"); fig.text(.06,.82,"公司整体完成率55%，4月和8月超过70%，2月和6月较低未过半",color=WHITE,fontsize=14,fontproperties=fp,va="top"); fig.text(.045,.045,"*注：数据来源于公司销售系统，统计日期截至2022.08.31",color="#D9D9D9",fontsize=8,fontproperties=fp)
    out=base_dir()/OUTPUT; fig.savefig(out,dpi=300,facecolor=BG); print(f"图片已保存：{out.resolve()}"); plt.show()
if __name__=="__main__": main()
