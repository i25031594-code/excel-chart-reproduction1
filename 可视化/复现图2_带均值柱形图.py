"""复现图2：带均值柱形图。兼容 Jupyter 和普通 .py 文件。"""
from pathlib import Path
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.font_manager import FontProperties
from openpyxl import load_workbook

EXCEL_FILE = "第二章 图表(前15).xlsx"
SHEET = "2 带均值柱形图"
OUTPUT = "图2_带均值柱形图.png"
BG, WHITE, BLUE1, BLUE2, GOLD = "#1A1E43", "#FFFFFF", "#0070C0", "#00B0F0", "#FFC000"

def base_dir(): return Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
def excel_path():
    for p in (base_dir()/EXCEL_FILE, Path.cwd()/EXCEL_FILE, Path.cwd()/"upload"/EXCEL_FILE):
        if p.exists(): return p
    raise FileNotFoundError(f"未找到{EXCEL_FILE}，请将它与代码放在同一文件夹。")
def font():
    paths=["C:/Windows/Fonts/msyh.ttc","C:/Windows/Fonts/simhei.ttf","/System/Library/Fonts/PingFang.ttc","/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]
    for p in paths:
        if Path(p).exists(): return FontProperties(fname=p)
    return FontProperties(family=["Microsoft YaHei","SimHei","sans-serif"])
def read_data():
    wb=load_workbook(excel_path(),data_only=True,read_only=True); ws=wb[SHEET]
    names=[ws.cell(r,2).value for r in range(3,9)]; values=np.array([ws.cell(r,3).value for r in range(3,9)],float); wb.close()
    return names,values
def gradient_bars(ax,x,y,width=.31):
    bars=ax.bar(x,y,width,color="none",edgecolor="none",zorder=3)
    cmap=LinearSegmentedColormap.from_list("blue",[BLUE2,BLUE1]); grad=np.linspace(0,1,400)[:,None]
    for b in bars:
        im=ax.imshow(grad,extent=[b.get_x(),b.get_x()+b.get_width(),0,b.get_height()],origin="lower",aspect="auto",cmap=cmap,zorder=3); im.set_clip_path(b)
    return bars
def main():
    names,y=read_data(); avg=y.mean(); fp=font(); x=np.arange(len(names))
    fig=plt.figure(figsize=(11.56,7.92),dpi=150,facecolor=BG); ax=fig.add_axes([.06,.16,.88,.53],facecolor=BG)
    bars=gradient_bars(ax,x,y); ax.axhline(avg,color=GOLD,lw=1.6,zorder=4)
    ax.set_xlim(-.5,5.5); ax.set_ylim(0,4000); ax.set_xticks(x,names); ax.set_yticks([]); ax.tick_params(axis="x",length=0,colors=WHITE,pad=8)
    for t in ax.get_xticklabels(): t.set_fontproperties(fp); t.set_fontsize(9)
    for s in ax.spines.values(): s.set_visible(False)
    ax.spines["bottom"].set_visible(True); ax.spines["bottom"].set_color("#A7ABC0"); ax.spines["bottom"].set_alpha(.55)
    for b,v in zip(bars,y): ax.text(b.get_x()+b.get_width()/2,v+70,f"{int(v)}",ha="center",color=WHITE,fontsize=9,fontproperties=fp)
    ax.text(5.13,avg+95,f"平均值 {avg:.0f}",color=GOLD,fontsize=10,ha="right",fontproperties=fp)
    fig.text(.06,.92,"3月各区域销量分布",color=WHITE,fontsize=20,fontweight="bold",fontproperties=fp,va="top")
    fig.text(.06,.82,"东北销量最多占比总销量的22%，华南销量最低",color=WHITE,fontsize=14,fontproperties=fp,va="top")
    fig.text(.045,.045,"*注：数据来源于公司销售系统，统计日期截至2022.03.31",color="#D9D9D9",fontsize=8,fontproperties=fp)
    out=base_dir()/OUTPUT; fig.savefig(out,dpi=300,facecolor=BG); print(f"图片已保存：{out.resolve()}"); plt.show()
if __name__=="__main__": main()
