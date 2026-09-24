# -*- coding: utf-8 -*-
"""图28 复合柱形图：严格读取 Excel 数据，完整显示中文标题、标签、图例与脚注。"""
from pathlib import Path

import numpy as np
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.font_manager import FontProperties
from matplotlib.patches import Circle, Wedge, FancyBboxPatch, Rectangle
from openpyxl import load_workbook

EXCEL_FILE = "第二章 图表(后15).xlsx"
SHEET_NAME = "28 复合柱形图"
OUTPUT_FILE = "图28_复合柱形图.png"

BG = "#191E45"
WHITE = "#F7F8FC"
TEXT2 = "#D7DBEA"
MUTED = "#AEB5CE"
GRID = "#4B5277"
BLUE = "#087CC1"
BRIGHT_BLUE = "#1597E5"
CYAN = "#20A8B5"
LIGHT_BLUE = "#45A5D6"
PINK = "#EB4D6D"
YELLOW = "#FBC44F"
PURPLE = "#6936D9"
ORANGE = "#FF7A1A"


def load_chinese_fonts():
    """优先直接加载 Windows 中文字体文件，避免 Jupyter 中文丢失。"""
    regular_paths = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/msyhl.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
        Path("/System/Library/Fonts/PingFang.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    ]
    bold_paths = [
        Path("C:/Windows/Fonts/msyhbd.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("/System/Library/Fonts/PingFang.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"),
    ]

    regular = next((p for p in regular_paths if p.exists()), None)
    bold = next((p for p in bold_paths if p.exists()), regular)

    if regular is not None:
        font_manager.fontManager.addfont(str(regular))
        normal_prop = FontProperties(fname=str(regular))
        bold_prop = FontProperties(fname=str(bold or regular), weight="bold")
        plt.rcParams["font.family"] = normal_prop.get_name()
    else:
        preferred = ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC",
                     "Source Han Sans SC", "PingFang SC", "Arial Unicode MS"]
        available = {f.name for f in font_manager.fontManager.ttflist}
        selected = next((name for name in preferred if name in available), "DejaVu Sans")
        normal_prop = FontProperties(family=selected)
        bold_prop = FontProperties(family=selected, weight="bold")
        plt.rcParams["font.family"] = selected

    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.dpi"] = 150
    plt.rcParams["savefig.dpi"] = 300
    return normal_prop, bold_prop


FONT, FONT_BOLD = load_chinese_fonts()


def tx(owner, x, y, text, *, bold=False, **kwargs):
    """统一文字函数：所有中文都显式指定 FontProperties。"""
    kwargs["fontproperties"] = FONT_BOLD if bold else FONT
    return owner.text(x, y, text, **kwargs)


def find_excel_file():
    """兼容 Jupyter，不使用 __file__。"""
    candidates = [
        Path.cwd() / EXCEL_FILE,
        Path.cwd() / "data" / EXCEL_FILE,
        Path.cwd() / "upload" / EXCEL_FILE,
        Path.cwd().parent / EXCEL_FILE,
    ]
    for path in candidates:
        if path.exists():
            return path
    raise FileNotFoundError(
        f"找不到 {EXCEL_FILE}。请把 Excel 与 Jupyter/.py 文件放在同一文件夹。"
    )


def read_rows(columns, start_row, end_row):
    workbook = load_workbook(find_excel_file(), data_only=True, read_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        workbook.close()
        raise KeyError(f"找不到工作表：{SHEET_NAME}")
    ws = workbook[SHEET_NAME]
    rows = [[ws[f"{col}{row}"].value for col in columns]
            for row in range(start_row, end_row + 1)]
    workbook.close()
    return rows


def canvas(figsize=(10, 6.2)):
    fig, ax = plt.subplots(figsize=figsize, facecolor=BG)
    ax.set_facecolor(BG)
    return fig, ax


def header(fig, year, title, subtitle, metrics=()):
    tx(fig, 0.055, 0.925, year, color=WHITE, fontsize=21, bold=True)
    tx(fig, 0.055, 0.850, title, color=WHITE, fontsize=17, bold=True)
    if subtitle:
        tx(fig, 0.055, 0.805, subtitle, color=TEXT2, fontsize=10.5)
    if metrics:
        right = 0.94
        step = min(0.145, 0.50 / len(metrics))
        start = right - step * (len(metrics) - 1)
        for i, (value, label) in enumerate(metrics):
            x = start + i * step
            tx(fig, x, 0.925, value, color=WHITE, fontsize=13.5, ha="center")
            tx(fig, x, 0.894, label, color=MUTED, fontsize=7.3, ha="center")


def footer(fig, date="2022.06.30", source="注：数据来源于公司人力资源系统"):
    tx(fig, 0.055, 0.050, "*", color=WHITE, fontsize=9)
    tx(fig, 0.075, 0.050, source, color=MUTED, fontsize=7.5)
    tx(fig, 0.945, 0.050, date, color=TEXT2, fontsize=8, ha="right")


def style_category_axis(ax):
    for tick in ax.get_xticklabels() + ax.get_yticklabels():
        tick.set_fontproperties(FONT)
        tick.set_color(WHITE)


def finish(fig):
    output = Path.cwd() / OUTPUT_FILE
    fig.savefig(output, facecolor=BG, dpi=300)
    print(f"图片已保存：{output}")
    plt.show()

def create_chart():
    rows=[r for r in read_rows(["B","C","D"],3,17) if r[0] is not None]
    months=[str(r[0]) for r in rows]; sales=np.array([float(r[1]) for r in rows]); qtot=np.array([float(r[2]) for r in rows])
    fig,ax=canvas((11,6.2))
    header(fig,"2021","2021年各月化妆品销量走势","2021年第三季度销量最多9673，9月单月销量最大3621",[("2021","统计年度"),("9673","最高季度销量"),("3621","最高单月销量")])
    footer(fig,date="2021.12.31",source="注：数据来源于公司销售系统，统计日期截至2021.12.31")
    ax.set_position([.08,.16,.84,.54]); x=np.arange(12); cols=["#586BEC",ORANGE,"#FFB000","#B83C82"]
    for q,col in enumerate(cols):
        idx=np.arange(q*3,q*3+3)
        ax.bar(idx,qtot[idx],.48,color=col,alpha=.18,zorder=1)
        ax.bar(idx,sales[idx],.48,color=col,zorder=3,label=f"第{q+1}季度")
        tx(ax,q*3+1,qtot[q*3]+160,str(int(qtot[q*3])),color=WHITE,fontsize=8.5,ha="center")
    for xx,v in zip(x,sales): tx(ax,xx,v+70,str(int(v)),color=WHITE,fontsize=7.2,ha="center")
    ax.set_xticks(x,months); style_category_axis(ax); ax.set_ylim(0,10800)
    ax.tick_params(axis="y",left=False,labelleft=False); ax.spines[:].set_visible(False); ax.axhline(0,color=GRID); ax.set_facecolor(BG)
    leg=ax.legend(loc="upper center",ncol=4,frameon=False,bbox_to_anchor=(.5,1.04),labelcolor=WHITE,fontsize=7.2)
    for t in leg.get_texts(): t.set_fontproperties(FONT)
    finish(fig)


if __name__ == "__main__": create_chart()

