# -*- coding: utf-8 -*-
"""图24 目标柱形图：按照参考图复现。"""

from pathlib import Path

import numpy as np
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.font_manager import FontProperties
from matplotlib.patches import Rectangle, Patch
from openpyxl import load_workbook


EXCEL_FILE = "第二章 图表(后15).xlsx"
SHEET_NAME = "24 目标柱形图"
OUTPUT_FILE = "图24_目标柱形图.png"

BG = "#191E45"
WHITE = "#F7F8FC"
TEXT2 = "#D7DBEA"
MUTED = "#AEB5CE"
GRID = "#27345E"
BLUE = "#087CC1"
TARGET = "#9FA8C1"


def load_chinese_fonts():
    """加载中文字体。"""

    regular_paths = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/msyhl.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
        Path("/System/Library/Fonts/PingFang.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    ]

    bold_paths = [
        Path("C:/Windows/Fonts/msyhbd.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("/System/Library/Fonts/PingFang.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"),
    ]

    regular = next(
        (path for path in regular_paths if path.exists()),
        None
    )

    bold = next(
        (path for path in bold_paths if path.exists()),
        regular
    )

    if regular is not None:
        font_manager.fontManager.addfont(str(regular))

        normal_prop = FontProperties(fname=str(regular))
        bold_prop = FontProperties(
            fname=str(bold or regular),
            weight="bold"
        )

        plt.rcParams["font.family"] = normal_prop.get_name()

    else:
        preferred = [
            "Microsoft YaHei",
            "SimHei",
            "Noto Sans CJK SC",
            "Source Han Sans SC",
            "PingFang SC",
            "Arial Unicode MS",
        ]

        available = {
            font.name
            for font in font_manager.fontManager.ttflist
        }

        selected = next(
            (
                name
                for name in preferred
                if name in available
            ),
            "DejaVu Sans"
        )

        normal_prop = FontProperties(family=selected)
        bold_prop = FontProperties(
            family=selected,
            weight="bold"
        )

        plt.rcParams["font.family"] = selected

    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.dpi"] = 150
    plt.rcParams["savefig.dpi"] = 300

    return normal_prop, bold_prop


FONT, FONT_BOLD = load_chinese_fonts()


def tx(owner, x, y, text, *, bold=False, **kwargs):
    """统一设置中文字体。"""

    kwargs["fontproperties"] = (
        FONT_BOLD if bold else FONT
    )

    return owner.text(x, y, text, **kwargs)


def find_excel_file():
    """查找 Excel 文件。"""

    candidates = [
        Path.cwd() / EXCEL_FILE,
        Path.cwd() / "data" / EXCEL_FILE,
        Path.cwd() / "upload" / EXCEL_FILE,
        Path.cwd().parent / EXCEL_FILE,
    ]

    for path in candidates:
        if path.exists():
            return path

    raise FileNotFoundError(
        f"找不到 {EXCEL_FILE}。"
        "请把 Excel 文件与 Python 文件放在同一个文件夹。"
    )


def read_rows(columns, start_row, end_row):
    """读取 Excel 数据。"""

    workbook = load_workbook(
        find_excel_file(),
        data_only=True,
        read_only=True
    )

    if SHEET_NAME not in workbook.sheetnames:
        workbook.close()
        raise KeyError(f"找不到工作表：{SHEET_NAME}")

    worksheet = workbook[SHEET_NAME]

    rows = [
        [
            worksheet[f"{column}{row}"].value
            for column in columns
        ]
        for row in range(start_row, end_row + 1)
    ]

    workbook.close()

    return rows


def create_canvas(figsize=(10, 6.2)):
    """创建深色背景画布。"""

    fig, ax = plt.subplots(
        figsize=figsize,
        facecolor=BG
    )

    ax.set_facecolor(BG)

    return fig, ax


def draw_title(fig):
    """绘制标题和副标题。"""

    tx(
        fig,
        0.075,
        0.900,
        "2022年上半年各商品销量完成情况",
        color=WHITE,
        fontsize=21,
        bold=True
    )

    tx(
        fig,
        0.075,
        0.850,
        "防晒整体销量最好，达到856，面霜远超目标，超额完成30%",
        color=TEXT2,
        fontsize=12
    )


def draw_footer(fig):
    """绘制脚注。"""

    tx(
        fig,
        0.075,
        0.055,
        "*注：数据来源于公司销售系统，统计日期截至2022.06.30",
        color=MUTED,
        fontsize=8.5
    )


def style_axis(ax):
    """设置坐标轴样式。"""

    ax.set_facecolor(BG)

    ax.tick_params(
        axis="y",
        left=False,
        labelleft=False
    )

    ax.tick_params(
        axis="x",
        length=0,
        pad=8
    )

    for label in ax.get_xticklabels():
        label.set_fontproperties(FONT)
        label.set_color(WHITE)
        label.set_fontsize(10)

    for spine in ax.spines.values():
        spine.set_visible(False)

    ax.axhline(
        0,
        color=GRID,
        linewidth=1.0,
        zorder=1
    )


def create_chart():
    """创建目标销量与实际销量对比图。"""

    # B列：商品名称
    # C列：实际销量
    # D列：目标销量
    rows = read_rows(
        columns=["B", "C", "D"],
        start_row=3,
        end_row=8
    )

    labels = [
        str(row[0]).strip()
        for row in rows
    ]

    actual = np.array(
        [
            float(row[1])
            for row in rows
        ]
    )

    target = np.array(
        [
            float(row[2])
            for row in rows
        ]
    )

    fig, ax = create_canvas()

    draw_title(fig)
    draw_footer(fig)

    # 设置图表位置
    ax.set_position([
        0.075,
        0.185,
        0.84,
        0.54
    ])

    x = np.arange(len(labels))

    # 实际销量柱宽
    actual_width = 0.18

    # 目标销量边框宽度
    target_width = 0.25

    # 目标销量：空心边框柱
    for position, target_value in zip(x, target):
        target_bar = Rectangle(
            (
                position - target_width / 2,
                0
            ),
            target_width,
            target_value,
            facecolor="none",
            edgecolor=TARGET,
            linewidth=1.5,
            zorder=2
        )

        ax.add_patch(target_bar)

    # 实际销量：蓝色实心柱
    bars = ax.bar(
        x,
        actual,
        width=actual_width,
        color=BLUE,
        edgecolor="none",
        zorder=3,
        label="实际销量"
    )

    # 添加实际销量数字
    for bar, value in zip(bars, actual):
        tx(
            ax,
            bar.get_x() + bar.get_width() / 2,
            value + 25,
            f"{int(value)}",
            color=WHITE,
            fontsize=10,
            ha="center",
            va="bottom"
        )

    ax.set_xticks(x)
    ax.set_xticklabels(labels)

    ax.set_ylim(
        0,
        max(max(actual), max(target)) * 1.18
    )

    style_axis(ax)

    # 创建图例
    legend_handles = [
        Patch(
            facecolor="none",
            edgecolor=TARGET,
            linewidth=1.5,
            label="目标销量"
        ),
        Patch(
            facecolor=BLUE,
            edgecolor="none",
            label="实际销量"
        )
    ]

    legend = ax.legend(
        handles=legend_handles,
        loc="upper left",
        bbox_to_anchor=(0.02, 1.18),
        ncol=2,
        frameon=False,
        handlelength=0.9,
        handleheight=0.8,
        columnspacing=1.0,
        borderaxespad=0,
        fontsize=10
    )

    for text in legend.get_texts():
        text.set_fontproperties(FONT)
        text.set_color(WHITE)

    fig.savefig(
        Path.cwd() / OUTPUT_FILE,
        facecolor=BG,
        dpi=300,
        bbox_inches="tight"
    )

    print(f"图片已保存：{Path.cwd() / OUTPUT_FILE}")

    plt.show()


if __name__ == "__main__":
    create_chart()
