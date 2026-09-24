"""复现图5：季度销售额与利润额层叠式对比柱形图。"""
from pathlib import Path
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from openpyxl import load_workbook
EXCEL_FILE="第二章 图表(前15).xlsx"; SHEET="5 层叠柱形图"; OUTPUT="图5_层叠柱形图.png"; BG="#1A1E43"; WHITE="#FFFFFF"; BLUE="#0070C0"; RED="#E74E69"
def base_dir(): return Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
def excel_path():
    for p in (base_dir()/EXCEL_FILE,Path.cwd()/EXCEL_FILE,Path.cwd()/"upload"/EXCEL_FILE):
        if p.exists(): return p
    raise FileNotFoundError(f"未找到{EXCEL_FILE}")
def font():
    for p in ["C:/Windows/Fonts/msyh.ttc","C:/Windows/Fonts/simhei.ttf","/System/Library/Fonts/PingFang.ttc","/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]:
        if Path(p).exists(): return FontProperties(fname=p)
    return FontProperties(family=["Microsoft YaHei","SimHei","sans-serif"])
def read_data():
    wb=load_workbook(excel_path(),data_only=True,read_only=True); ws=wb[SHEET]; q=[ws.cell(r,2).value for r in range(3,9)]; s=np.array([ws.cell(r,3).value for r in range(3,9)],float); p=np.array([ws.cell(r,4).value for r in range(3,9)],float); wb.close(); return q,s,p
def main():
    q,s,p=read_data(); fp=font(); x=np.arange(6); fig=plt.figure(figsize=(11.56,7.72),dpi=150,facecolor=BG); ax=fig.add_axes([.09,.15,.82,.52],facecolor=BG)
    b1=ax.bar(x-.10,s,.26,color=BLUE,edgecolor="none",zorder=2); b2=ax.bar(x+.10,p,.26,color=RED,edgecolor="none",zorder=3); ax.set_ylim(0,5600); ax.set_xticks(x,q); ax.set_yticks([]); ax.tick_params(axis="x",length=0,colors=WHITE,pad=8)
    for t in ax.get_xticklabels(): t.set_fontproperties(fp); t.set_fontsize(9)
    for sp in ax.spines.values(): sp.set_visible(False)
    for bars,vals in [(b1,s),(b2,p)]:
        for b,v in zip(bars,vals): ax.text(b.get_x()+b.get_width()/2,v+90,f"{int(v)}",ha="center",color=WHITE,fontsize=9,fontproperties=fp)
    ax.text(5.15,2600,"销售额",color=WHITE,fontsize=9,fontproperties=fp,bbox=dict(boxstyle="square,pad=.45",fc="none",ec=BLUE,lw=1))
    ax.text(5.15,1850,"利润额",color=WHITE,fontsize=9,fontproperties=fp,bbox=dict(boxstyle="square,pad=.45",fc="none",ec=RED,lw=1))
    fig.text(.055,.92,"2021年至今季度销售额(万)和利润额(万)",color=WHITE,fontsize=20,fontweight="bold",fontproperties=fp,va="top")
    fig.text(.055,.82,"2022年第二季度销售额首次出现下降，降幅达到15%",color=WHITE,fontsize=14,fontproperties=fp,va="top")
    fig.text(.045,.045,"*注：数据来源于公司销售系统，统计日期截至2022.06.30",color="#D9D9D9",fontsize=8,fontproperties=fp)
    out=base_dir()/OUTPUT; fig.savefig(out,dpi=300,facecolor=BG); print(f"图片已保存：{out.resolve()}"); plt.show()
if __name__=="__main__": main()
