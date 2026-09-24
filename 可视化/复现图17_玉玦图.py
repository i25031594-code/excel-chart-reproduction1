# -*- coding: utf-8 -*-
"""图17 玉玦图：图例放置在图形右侧，避免与环形图重合。"""

from pathlib import Path

import numpy as np
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.font_manager import FontProperties
from matplotlib.patches import Wedge, Patch
from openpyxl import load_workbook


EXCEL_FILE = "第二章 图表(后15).xlsx"
SHEET_NAME = "17 玉玦图"
OUTPUT_FILE = "图17_玉玦图.png"

BG = "#191E45"
WHITE = "#F7F8FC"
TEXT2 = "#D7DBEA"
MUTED = "#AEB5CE"

CYAN = "#20A8B5"
LIGHT_BLUE = "#45A5D6"
PINK = "#EB4D6D"
YELLOW = "#FBC44F"


def load_chinese_fonts():
    """加载中文字体。"""

    regular_paths = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/msyhl.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
        Path("/System/Library/Fonts/PingFang.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    ]

    bold_paths = [
        Path("C:/Windows/Fonts/msyhbd.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("/System/Library/Fonts/PingFang.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"),
    ]

    regular = next((p for p in regular_paths if p.exists()), None)
    bold = next((p for p in bold_paths if p.exists()), regular)

    if regular is not None:
        font_manager.fontManager.addfont(str(regular))

        normal_prop = FontProperties(fname=str(regular))
        bold_prop = FontProperties(
            fname=str(bold or regular),
            weight="bold"
        )

        plt.rcParams["font.family"] = normal_prop.get_name()

    else:
        preferred = [
            "Microsoft YaHei",
            "SimHei",
            "Noto Sans CJK SC",
            "Source Han Sans SC",
            "PingFang SC",
            "Arial Unicode MS",
        ]

        available = {
            font.name for font in font_manager.fontManager.ttflist
        }

        selected = next(
            (name for name in preferred if name in available),
            "DejaVu Sans"
        )

        normal_prop = FontProperties(family=selected)
        bold_prop = FontProperties(
            family=selected,
            weight="bold"
        )

        plt.rcParams["font.family"] = selected

    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.dpi"] = 150
    plt.rcParams["savefig.dpi"] = 300

    return normal_prop, bold_prop


FONT, FONT_BOLD = load_chinese_fonts()


def tx(owner, x, y, text, *, bold=False, **kwargs):
    """统一设置中文字体。"""

    kwargs["fontproperties"] = FONT_BOLD if bold else FONT
    return owner.text(x, y, text, **kwargs)


def find_excel_file():
    """查找 Excel 文件。"""

    candidates = [
        Path.cwd() / EXCEL_FILE,
        Path.cwd() / "data" / EXCEL_FILE,
        Path.cwd() / "upload" / EXCEL_FILE,
        Path.cwd().parent / EXCEL_FILE,
    ]

    for path in candidates:
        if path.exists():
            return path

    raise FileNotFoundError(
        f"找不到 {EXCEL_FILE}。"
        "请把 Excel 文件与 Python 文件放在同一个文件夹。"
    )


def read_rows(columns, start_row, end_row):
    """读取 Excel 数据。"""

    workbook = load_workbook(
        find_excel_file(),
        data_only=True,
        read_only=True
    )

    if SHEET_NAME not in workbook.sheetnames:
        workbook.close()
        raise KeyError(f"找不到工作表：{SHEET_NAME}")

    worksheet = workbook[SHEET_NAME]

    rows = [
        [
            worksheet[f"{column}{row}"].value
            for column in columns
        ]
        for row in range(start_row, end_row + 1)
    ]

    workbook.close()
    return rows


def canvas(figsize=(10, 6.2)):
    """创建画布。"""

    fig, ax = plt.subplots(
        figsize=figsize,
        facecolor=BG
    )

    ax.set_facecolor(BG)

    return fig, ax


def header(fig, year, title, subtitle, metrics=()):
    """绘制标题和右上角指标。"""

    tx(
        fig,
        0.055,
        0.925,
        year,
        color=WHITE,
        fontsize=21,
        bold=True
    )

    tx(
        fig,
        0.055,
        0.850,
        title,
        color=WHITE,
        fontsize=17,
        bold=True
    )

    if subtitle:
        tx(
            fig,
            0.055,
            0.805,
            subtitle,
            color=TEXT2,
            fontsize=10.5
        )

    if metrics:
        right = 0.94
        step = min(0.145, 0.50 / len(metrics))
        start = right - step * (len(metrics) - 1)

        for index, (value, label) in enumerate(metrics):
            x = start + index * step

            tx(
                fig,
                x,
                0.925,
                value,
                color=WHITE,
                fontsize=13.5,
                ha="center"
            )

            tx(
                fig,
                x,
                0.894,
                label,
                color=MUTED,
                fontsize=7.3,
                ha="center"
            )


def footer(
    fig,
    date="2022.06.30",
    source="注：数据来源于公司人力资源系统"
):
    """绘制脚注。"""

    tx(
        fig,
        0.055,
        0.050,
        "*",
        color=WHITE,
        fontsize=9
    )

    tx(
        fig,
        0.075,
        0.050,
        source,
        color=MUTED,
        fontsize=7.5
    )

    tx(
        fig,
        0.945,
        0.050,
        date,
        color=TEXT2,
        fontsize=8,
        ha="right"
    )


def finish(fig):
    """保存并显示图片。"""

    output = Path.cwd() / OUTPUT_FILE

    fig.savefig(
        output,
        facecolor=BG,
        dpi=300,
        bbox_inches="tight"
    )

    print(f"图片已保存：{output}")
    plt.show()


def create_chart():
    """创建年龄分布环形图。"""

    rows = read_rows(
        columns=["B", "C"],
        start_row=3,
        end_row=6
    )

    labels = [
        str(row[0]).strip()
        for row in rows
    ][::-1]

    values = np.array(
        [
            float(row[1])
            for row in rows
        ]
    )[::-1]

    colors = [
        PINK,
        YELLOW,
        CYAN,
        LIGHT_BLUE
    ]

    fig, ax = canvas()

    header(
        fig,
        "2022",
        "2022年上半年年龄分布",
        "公司平均年龄32.5，23-30员工比例最高",
        [
            ("32.5", "公司平均年龄"),
            ("23-30", "员工集中年龄")
        ]
    )

    footer(fig)

    # 缩小环形图区域，为右侧图例留出空间
    ax.set_position([
        0.16,
        0.12,
        0.55,
        0.66
    ])

    ax.set_aspect("equal")
    ax.axis("off")

    start_angle = 180

    for index, (label, value, color) in enumerate(
        zip(labels, values, colors)
    ):
        radius = 1.0 - index * 0.16

        value_angle = 360 * (value / 0.5)

        # 彩色部分
        ax.add_patch(
            Wedge(
                center=(0, 0),
                r=radius,
                theta1=start_angle,
                theta2=start_angle + value_angle,
                width=0.12,
                facecolor=color,
                edgecolor="none"
            )
        )

        # 背景部分
        ax.add_patch(
            Wedge(
                center=(0, 0),
                r=radius,
                theta1=start_angle + value_angle,
                theta2=start_angle + 360,
                width=0.12,
                facecolor="#30365E",
                edgecolor="none",
                alpha=0.75
            )
        )

    # 中心文字
    tx(
        ax,
        0,
        0,
        "年龄\n分布",
        color=WHITE,
        fontsize=11,
        ha="center",
        va="center"
    )

    # 设置图形范围
    ax.set_xlim(-1.25, 1.25)
    ax.set_ylim(-1.16, 1.25)

    # 创建右侧独立图例
    legend_handles = [
        Patch(
            facecolor=color,
            edgecolor="none",
            label=f"{label}    {value:.1%}"
        )
        for label, value, color in zip(
            labels,
            values,
            colors
        )
    ]

    legend = fig.legend(
        handles=legend_handles,
        loc="center right",
        bbox_to_anchor=(0.96, 0.48),
        frameon=False,
        fontsize=9,
        labelcolor=WHITE,
        handlelength=1.2,
        handleheight=1.2,
        borderaxespad=0.0
    )

    # 确保图例中文字体正常显示
    for text in legend.get_texts():
        text.set_fontproperties(FONT)
        text.set_color(WHITE)

    finish(fig)


if __name__ == "__main__":
    create_chart()
