"""复现图8：区域销量及同比百分比水平堆叠图。"""
from pathlib import Path
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from openpyxl import load_workbook
EXCEL_FILE="第二章 图表(前15).xlsx"; SHEET="8 数值百分比"; OUTPUT="图8_数值百分比.png"; BG="#1A1E43"; WHITE="#FFFFFF"
def base_dir(): return Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
def excel_path():
    for p in (base_dir()/EXCEL_FILE,Path.cwd()/EXCEL_FILE,Path.cwd()/"upload"/EXCEL_FILE):
        if p.exists(): return p
    raise FileNotFoundError(f"未找到{EXCEL_FILE}")
def font():
    for p in ["C:/Windows/Fonts/msyh.ttc","C:/Windows/Fonts/simhei.ttf","/System/Library/Fonts/PingFang.ttc","/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]:
        if Path(p).exists(): return FontProperties(fname=p)
    return FontProperties(family=["Microsoft YaHei","SimHei","sans-serif"])
def read_data():
    wb=load_workbook(excel_path(),data_only=True,read_only=True); ws=wb[SHEET]; n=[ws.cell(r,2).value for r in range(3,9)]; v=np.array([ws.cell(r,3).value for r in range(3,9)],float); p=np.array([ws.cell(r,6).value for r in range(3,9)],float)*100; wb.close(); return n,v,p
def main():
    n,v,p=read_data(); n=n[::-1]; v=v[::-1]; p=p[::-1]; fp=font(); y=np.arange(6)[::-1]; mx=v.max(); spacer=mx-v; redw=mx/4
    fig=plt.figure(figsize=(11.56,7.72),dpi=150,facecolor=BG); ax=fig.add_axes([.12,.14,.77,.57],facecolor=BG)
    ax.barh(y,v,.66,color="#09387E"); ax.barh(y,spacer,.66,left=v,color="#82ADD7"); ax.barh(y,np.full(6,redw),.66,left=mx,color="#9B3D4F")
    ax.set_xlim(0,mx+redw*1.12); ax.set_yticks(y,n); ax.set_xticks([]); ax.tick_params(axis="y",length=0,colors=WHITE,pad=8)
    for t in ax.get_yticklabels(): t.set_fontproperties(fp); t.set_fontsize(9)
    for s in ax.spines.values(): s.set_visible(False)
    for yi,val,pct in zip(y,v,p):
        ax.text(val-35,yi,f"{int(val)}",ha="right",va="center",color=WHITE,fontsize=9,fontproperties=fp)
        ax.text(mx+redw/2,yi,f"{pct:.1f}%",ha="center",va="center",color=WHITE,fontsize=9,fontproperties=fp)
    fig.text(.06,.92,"2021年各区域销量及同比情况",color=WHITE,fontsize=20,fontweight="bold",fontproperties=fp,va="top")
    fig.text(.06,.82,"各区域商品销量同比去年均有下降，其中华南下降最多，同比下降20.8%",color=WHITE,fontsize=14,fontproperties=fp,va="top")
    fig.text(.045,.045,"*注：数据来源于公司销售系统，统计日期截至2022.01.01",color="#D9D9D9",fontsize=8,fontproperties=fp)
    out=base_dir()/OUTPUT; fig.savefig(out,dpi=300,facecolor=BG); print(f"图片已保存：{out.resolve()}"); plt.show()
if __name__=="__main__": main()
