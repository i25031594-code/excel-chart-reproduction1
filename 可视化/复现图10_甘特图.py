"""复现图10：2022年化妆品类目采购项目甘特图。"""
from pathlib import Path
from datetime import datetime, timedelta
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from matplotlib.font_manager import FontProperties
from openpyxl import load_workbook
EXCEL_FILE="第二章 图表(前15).xlsx"; SHEET="10 甘特图"; OUTPUT="图10_甘特图.png"; BG="#1A1E43"; WHITE="#FFFFFF"; BLUE="#0070C0"; CYAN="#00B0F0"
def base_dir(): return Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
def excel_path():
    for p in (base_dir()/EXCEL_FILE,Path.cwd()/EXCEL_FILE,Path.cwd()/"upload"/EXCEL_FILE):
        if p.exists(): return p
    raise FileNotFoundError(f"未找到{EXCEL_FILE}")
def font():
    for p in ["C:/Windows/Fonts/msyh.ttc","C:/Windows/Fonts/simhei.ttf","/System/Library/Fonts/PingFang.ttc","/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]:
        if Path(p).exists(): return FontProperties(fname=p)
    return FontProperties(family=["Microsoft YaHei","SimHei","sans-serif"])
def read_data():
    wb=load_workbook(excel_path(),data_only=True,read_only=True); ws=wb[SHEET]
    names=[ws.cell(r,2).value for r in range(4,11)]; starts=[ws.cell(r,3).value for r in range(4,11)]; days=np.array([ws.cell(r,4).value for r in range(4,11)],float); rates=np.array([ws.cell(r,5).value for r in range(4,11)],float); wb.close(); return names,starts,days,rates
def date_label(x,pos=None):
    d=mdates.num2date(x); return f"{d.month}/{d.day}/{d.year}"
def main():
    names,starts,days,rates=read_data(); fp=font(); left=mdates.date2num(starts); y=np.arange(7)[::-1]
    fig=plt.figure(figsize=(12.32,7.72),dpi=150,facecolor=BG); ax=fig.add_axes([.16,.13,.75,.63],facecolor=BG)
    ax.barh(y,days,left=left,height=.48,color=BLUE,edgecolor="none",zorder=2); ax.barh(y,days*rates,left=left,height=.48,color=CYAN,edgecolor="none",zorder=3)
    for yi,l,w,r in zip(y,left,days,rates): ax.text(l+w*r/2,yi,f"{r:.0%}",ha="center",va="center",color=WHITE,fontsize=9,fontproperties=fp,zorder=4)
    start_date=datetime(2022,3,1); tick_dates=[start_date+timedelta(days=15*i) for i in range(8)]
    ax.set_yticks(y,names); ax.set_xlim(mdates.date2num(tick_dates[0]),mdates.date2num(tick_dates[-1])); ax.set_xticks(mdates.date2num(tick_dates)); ax.xaxis.set_major_formatter(date_label); ax.xaxis.tick_top(); ax.tick_params(axis="both",length=0,colors=WHITE,pad=8)
    for t in list(ax.get_xticklabels())+list(ax.get_yticklabels()): t.set_fontproperties(fp); t.set_fontsize(9)
    ax.set_axisbelow(True); ax.grid(axis="x",color="white",alpha=.25,linestyle=(0,(5,4)),lw=.8)
    for s in ax.spines.values(): s.set_visible(False)
    fig.text(.5,.94,"2022年化妆品类目采购项目进度",ha="center",va="top",color=WHITE,fontsize=20,fontweight="bold",fontproperties=fp)
    out=base_dir()/OUTPUT; fig.savefig(out,dpi=300,facecolor=BG); print(f"图片已保存：{out.resolve()}"); plt.show()
if __name__=="__main__": main()
