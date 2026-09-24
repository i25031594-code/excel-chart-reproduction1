# -*- coding: utf-8 -*-
"""图23 柱形折线图：严格读取 Excel 数据，完整显示中文标题、标签、图例与脚注。"""
from pathlib import Path

import numpy as np
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.font_manager import FontProperties
from matplotlib.patches import Circle, Wedge, FancyBboxPatch, Rectangle
from openpyxl import load_workbook

EXCEL_FILE = "第二章 图表(后15).xlsx"
SHEET_NAME = "23 柱形折线图"
OUTPUT_FILE = "图23_柱形折线图.png"

BG = "#191E45"
WHITE = "#F7F8FC"
TEXT2 = "#D7DBEA"
MUTED = "#AEB5CE"
GRID = "#4B5277"
BLUE = "#087CC1"
BRIGHT_BLUE = "#1597E5"
CYAN = "#20A8B5"
LIGHT_BLUE = "#45A5D6"
PINK = "#EB4D6D"
YELLOW = "#FBC44F"
PURPLE = "#6936D9"
ORANGE = "#FF7A1A"


def load_chinese_fonts():
    """优先直接加载 Windows 中文字体文件，避免 Jupyter 中文丢失。"""
    regular_paths = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/msyhl.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
        Path("/System/Library/Fonts/PingFang.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    ]
    bold_paths = [
        Path("C:/Windows/Fonts/msyhbd.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("/System/Library/Fonts/PingFang.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"),
    ]

    regular = next((p for p in regular_paths if p.exists()), None)
    bold = next((p for p in bold_paths if p.exists()), regular)

    if regular is not None:
        font_manager.fontManager.addfont(str(regular))
        normal_prop = FontProperties(fname=str(regular))
        bold_prop = FontProperties(fname=str(bold or regular), weight="bold")
        plt.rcParams["font.family"] = normal_prop.get_name()
    else:
        preferred = ["Microsoft YaHei", "SimHei", "Noto Sans CJK SC",
                     "Source Han Sans SC", "PingFang SC", "Arial Unicode MS"]
        available = {f.name for f in font_manager.fontManager.ttflist}
        selected = next((name for name in preferred if name in available), "DejaVu Sans")
        normal_prop = FontProperties(family=selected)
        bold_prop = FontProperties(family=selected, weight="bold")
        plt.rcParams["font.family"] = selected

    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.dpi"] = 150
    plt.rcParams["savefig.dpi"] = 300
    return normal_prop, bold_prop


FONT, FONT_BOLD = load_chinese_fonts()


def tx(owner, x, y, text, *, bold=False, **kwargs):
    """统一文字函数：所有中文都显式指定 FontProperties。"""
    kwargs["fontproperties"] = FONT_BOLD if bold else FONT
    return owner.text(x, y, text, **kwargs)


def find_excel_file():
    """兼容 Jupyter，不使用 __file__。"""
    candidates = [
        Path.cwd() / EXCEL_FILE,
        Path.cwd() / "data" / EXCEL_FILE,
        Path.cwd() / "upload" / EXCEL_FILE,
        Path.cwd().parent / EXCEL_FILE,
    ]
    for path in candidates:
        if path.exists():
            return path
    raise FileNotFoundError(
        f"找不到 {EXCEL_FILE}。请把 Excel 与 Jupyter/.py 文件放在同一文件夹。"
    )


def read_rows(columns, start_row, end_row):
    workbook = load_workbook(find_excel_file(), data_only=True, read_only=True)
    if SHEET_NAME not in workbook.sheetnames:
        workbook.close()
        raise KeyError(f"找不到工作表：{SHEET_NAME}")
    ws = workbook[SHEET_NAME]
    rows = [[ws[f"{col}{row}"].value for col in columns]
            for row in range(start_row, end_row + 1)]
    workbook.close()
    return rows


def canvas(figsize=(10, 6.2)):
    fig, ax = plt.subplots(figsize=figsize, facecolor=BG)
    ax.set_facecolor(BG)
    return fig, ax


def header(fig, year, title, subtitle, metrics=()):
    tx(fig, 0.055, 0.925, year, color=WHITE, fontsize=21, bold=True)
    tx(fig, 0.055, 0.850, title, color=WHITE, fontsize=17, bold=True)
    if subtitle:
        tx(fig, 0.055, 0.805, subtitle, color=TEXT2, fontsize=10.5)
    if metrics:
        right = 0.94
        step = min(0.145, 0.50 / len(metrics))
        start = right - step * (len(metrics) - 1)
        for i, (value, label) in enumerate(metrics):
            x = start + i * step
            tx(fig, x, 0.925, value, color=WHITE, fontsize=13.5, ha="center")
            tx(fig, x, 0.894, label, color=MUTED, fontsize=7.3, ha="center")


def footer(fig, date="2022.06.30", source="注：数据来源于公司人力资源系统"):
    tx(fig, 0.055, 0.050, "*", color=WHITE, fontsize=9)
    tx(fig, 0.075, 0.050, source, color=MUTED, fontsize=7.5)
    tx(fig, 0.945, 0.050, date, color=TEXT2, fontsize=8, ha="right")


def style_category_axis(ax):
    for tick in ax.get_xticklabels() + ax.get_yticklabels():
        tick.set_fontproperties(FONT)
        tick.set_color(WHITE)


def finish(fig):
    output = Path.cwd() / OUTPUT_FILE
    fig.savefig(output, facecolor=BG, dpi=300)
    print(f"图片已保存：{output}")
    plt.show()

def create_chart():
    rows=read_rows(["B","C","D"],3,8)
    years=[str(r[0]) for r in rows]; sales=np.array([float(r[1]) for r in rows]); growth=np.array([float(r[2]) for r in rows])
    fig,ax=canvas()
    header(fig,"2022","近六年销售量及增长率","平台销量近6年持续增长，但近两年增长率有所放缓",[("6","统计年数")])
    footer(fig,source="注：数据来源于公司销售系统，统计日期截至2022.06.30")
    ax.set_position([.10,.16,.80,.54]); x=np.arange(len(years)); bars=ax.bar(x,sales,.34,color=BLUE,label="销售量",zorder=3)
    ax2=ax.twinx(); ax2.plot(x,growth,color=PINK,marker="o",ms=5.5,lw=2,label="同比",zorder=4)
    for b,v in zip(bars,sales): tx(ax,b.get_x()+b.get_width()/2,v+75,f"{int(v)}",color=WHITE,fontsize=8.5,ha="center")
    for xx,v in zip(x,growth): tx(ax2,xx,v+.022,f"{v:.0%}",color=WHITE,fontsize=8.5,ha="center")
    ax.set_ylim(0,4700); ax2.set_ylim(0,.48); ax.set_xticks(x,years); style_category_axis(ax)
    ax.tick_params(axis="y",left=False,labelleft=False); ax2.tick_params(right=False,labelright=False)
    ax.spines[:].set_visible(False); ax2.spines[:].set_visible(False); ax.axhline(0,color=GRID,lw=1)
    ax.set_facecolor(BG); ax2.set_facecolor("none")
    h1,l1=ax.get_legend_handles_labels(); h2,l2=ax2.get_legend_handles_labels()
    leg=ax.legend(h1+h2,l1+l2,loc="upper center",ncol=2,frameon=False,bbox_to_anchor=(.5,1.03),labelcolor=WHITE,fontsize=8)
    for t in leg.get_texts(): t.set_fontproperties(FONT)
    finish(fig)


if __name__ == "__main__": create_chart()

