"""复现图6：2022年与2021年区域销量蝴蝶图。"""
from pathlib import Path
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from openpyxl import load_workbook
EXCEL_FILE="第二章 图表(前15).xlsx"; SHEET="6 蝴蝶图"; OUTPUT="图6_销量蝴蝶图.png"; BG="#1A1E43"; WHITE="#FFFFFF"; BLUE="#0070C0"; RED="#E74E69"
def base_dir(): return Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
def excel_path():
    for p in (base_dir()/EXCEL_FILE,Path.cwd()/EXCEL_FILE,Path.cwd()/"upload"/EXCEL_FILE):
        if p.exists(): return p
    raise FileNotFoundError(f"未找到{EXCEL_FILE}")
def font():
    for p in ["C:/Windows/Fonts/msyh.ttc","C:/Windows/Fonts/simhei.ttf","/System/Library/Fonts/PingFang.ttc","/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]:
        if Path(p).exists(): return FontProperties(fname=p)
    return FontProperties(family=["Microsoft YaHei","SimHei","sans-serif"])
def read_data():
    wb=load_workbook(excel_path(),data_only=True,read_only=True); ws=wb[SHEET]; n=[ws.cell(r,2).value for r in range(3,8)]; a=np.array([ws.cell(r,4).value for r in range(3,8)],float); b=np.array([ws.cell(r,6).value for r in range(3,8)],float); wb.close(); return n,a,b
def main():
    n,a,b=read_data(); n=n[::-1]; a=a[::-1]; b=b[::-1]; fp=font(); y=np.arange(5)[::-1]
    fig=plt.figure(figsize=(11.56,7.72),dpi=150,facecolor=BG); ax=fig.add_axes([.12,.14,.76,.56],facecolor=BG)
    gap=230
    ax.barh(y,-a,.38,left=-gap,color=BLUE,edgecolor="none"); ax.barh(y,b,.38,left=gap,color=RED,edgecolor="none"); ax.set_xlim(-3000,3000); ax.set_ylim(-.8,5.0); ax.axis("off")
    for yi,name,l,r in zip(y,n,a,b):
        ax.text(0,yi,name,ha="center",va="center",color=WHITE,fontsize=9,fontproperties=fp)
        ax.text(-gap-l+35,yi,f"{int(l)}",ha="left",va="center",color=WHITE,fontsize=9,fontproperties=fp)
        ax.text(gap+r-35,yi,f"{int(r)}",ha="right",va="center",color=WHITE,fontsize=9,fontproperties=fp)
    ax.text(-1350,4.72,"2022",ha="center",color=BLUE,fontsize=11,fontproperties=fp); ax.text(1350,4.72,"2021",ha="center",color=RED,fontsize=11,fontproperties=fp)
    fig.text(.06,.92,"2022年上半年各区域对比去年销量",color=WHITE,fontsize=20,fontweight="bold",fontproperties=fp,va="top")
    fig.text(.06,.82,"2022年整体销量高于2021年，只有东北区域较2021有所下降",color=WHITE,fontsize=14,fontproperties=fp,va="top")
    fig.text(.045,.045,"*注：数据来源于公司销售系统，统计日期截至2022.06.30",color="#D9D9D9",fontsize=8,fontproperties=fp)
    out=base_dir()/OUTPUT; fig.savefig(out,dpi=300,facecolor=BG); print(f"图片已保存：{out.resolve()}"); plt.show()
if __name__=="__main__": main()
