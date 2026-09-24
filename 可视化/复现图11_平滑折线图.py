"""复现图11：月度销量平滑折线图。仅使用 NumPy 完成平滑，无需 SciPy。"""
from pathlib import Path
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from matplotlib.patches import FancyBboxPatch
from openpyxl import load_workbook
EXCEL_FILE="第二章 图表(前15).xlsx"; SHEET="11 平滑折线图"; OUTPUT="图11_平滑折线图.png"; BG="#1A1E43"; WHITE="#FFFFFF"; ORANGE="#E66B4C"
def base_dir(): return Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
def excel_path():
    for p in (base_dir()/EXCEL_FILE,Path.cwd()/EXCEL_FILE,Path.cwd()/"upload"/EXCEL_FILE):
        if p.exists(): return p
    raise FileNotFoundError(f"未找到{EXCEL_FILE}")
def font():
    for p in ["C:/Windows/Fonts/msyh.ttc","C:/Windows/Fonts/simhei.ttf","/System/Library/Fonts/PingFang.ttc","/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]:
        if Path(p).exists(): return FontProperties(fname=p)
    return FontProperties(family=["Microsoft YaHei","SimHei","sans-serif"])
def read_data():
    wb=load_workbook(excel_path(),data_only=True,read_only=True); ws=wb[SHEET]; m=[ws.cell(r,3).value for r in range(3,14)]; v=np.array([ws.cell(r,4).value for r in range(3,14)],float); wb.close(); return m,v
def hermite(x,y,points=600):
    slopes=np.empty_like(y); slopes[1:-1]=(y[2:]-y[:-2])/2; slopes[0]=y[1]-y[0]; slopes[-1]=y[-1]-y[-2]; xs=np.linspace(x[0],x[-1],points); ys=np.empty_like(xs)
    for j,z in enumerate(xs):
        i=min(int(np.floor(z)),len(x)-2); t=z-x[i]; h00=2*t**3-3*t**2+1; h10=t**3-2*t**2+t; h01=-2*t**3+3*t**2; h11=t**3-t**2; ys[j]=h00*y[i]+h10*slopes[i]+h01*y[i+1]+h11*slopes[i+1]
    return xs,ys
def main():
    months,v=read_data(); fp=font(); x=np.arange(len(v)); xs,ys=hermite(x,v); fig=plt.figure(figsize=(11.56,7.92),dpi=150,facecolor=BG); ax=fig.add_axes([.11,.22,.83,.50],facecolor=BG)
    ax.plot(xs,ys,color=ORANGE,lw=1.8); ax.vlines(8,0,v[8],color=ORANGE,linestyle=(0,(5,3)),lw=1.1,alpha=.85); ax.text(8,v[8]+110,f"{int(v[8])}",ha="center",color=WHITE,fontsize=9,fontproperties=fp)
    ax.set_xlim(-.5,10.5); ax.set_ylim(0,4000); ax.set_xticks(x,months); ax.set_yticks(np.arange(0,4001,1000),[f"{i:,}" if i else "0" for i in range(0,4001,1000)]); ax.tick_params(axis="x",colors=WHITE,pad=8); ax.tick_params(axis="y",colors=ORANGE,length=0,pad=8)
    for t in list(ax.get_xticklabels())+list(ax.get_yticklabels()): t.set_fontproperties(fp); t.set_fontsize(9)
    for s in ax.spines.values(): s.set_visible(False)
    ax.spines["bottom"].set_visible(True); ax.spines["bottom"].set_color("#A7ABC0"); ax.spines["bottom"].set_alpha(.75)
    ax.add_patch(FancyBboxPatch((-.35,-.25),7.85,.10,transform=ax.get_xaxis_transform(),clip_on=False,boxstyle="round,pad=.01,rounding_size=.04",fc="#67D0E6",ec="none")); ax.add_patch(FancyBboxPatch((7.65,-.25),3.05,.10,transform=ax.get_xaxis_transform(),clip_on=False,boxstyle="round,pad=.01,rounding_size=.04",fc="#F7C64D",ec="none"))
    ax.text(3.55,-.20,"2021",transform=ax.get_xaxis_transform(),ha="center",va="center",color=WHITE,fontsize=9,fontproperties=fp,clip_on=False); ax.text(9.15,-.20,"2022",transform=ax.get_xaxis_transform(),ha="center",va="center",color=WHITE,fontsize=9,fontproperties=fp,clip_on=False)
    fig.text(.055,.92,"化妆品品类月度销量走势",color=WHITE,fontsize=20,fontweight="bold",fontproperties=fp,va="top"); fig.text(.055,.82,"2022年销量迅速增加，1月最高，销量达到3782",color=WHITE,fontsize=14,fontproperties=fp,va="top"); fig.text(.045,.045,"*注：数据来源于公司销售系统，统计日期截至2022.03.31",color="#D9D9D9",fontsize=8,fontproperties=fp)
    out=base_dir()/OUTPUT; fig.savefig(out,dpi=300,facecolor=BG); print(f"图片已保存：{out.resolve()}"); plt.show()
if __name__=="__main__": main()
