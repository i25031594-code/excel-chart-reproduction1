"""复现图13：2022年上半年与2021年销量对比折线图。"""
from pathlib import Path
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from openpyxl import load_workbook
EXCEL_FILE="第二章 图表(前15).xlsx"; SHEET="13 对比折线图"; OUTPUT="图13_对比折线图.png"; BG="#1A1E43"; WHITE="#FFFFFF"; RED="#E74E69"; BLUE="#0070C0"
def base_dir(): return Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
def excel_path():
    for p in (base_dir()/EXCEL_FILE,Path.cwd()/EXCEL_FILE,Path.cwd()/"upload"/EXCEL_FILE):
        if p.exists(): return p
    raise FileNotFoundError(f"未找到{EXCEL_FILE}")
def font():
    for p in ["C:/Windows/Fonts/msyh.ttc","C:/Windows/Fonts/simhei.ttf","/System/Library/Fonts/PingFang.ttc","/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]:
        if Path(p).exists(): return FontProperties(fname=p)
    return FontProperties(family=["Microsoft YaHei","SimHei","sans-serif"])
def read_data():
    wb=load_workbook(excel_path(),data_only=True,read_only=True); ws=wb[SHEET]; m=[ws.cell(r,2).value for r in range(3,9)]; a=np.array([ws.cell(r,3).value for r in range(3,9)],float); b=np.array([ws.cell(r,4).value for r in range(3,9)],float); wb.close(); return m,a,b
def main():
    months,a,b=read_data(); fp=font(); x=np.arange(6); fig=plt.figure(figsize=(11.56,7.92),dpi=150,facecolor=BG); ax=fig.add_axes([.08,.15,.85,.53],facecolor=BG)
    ax.plot(x,a,color=RED,lw=2,marker="o",ms=6,label="2021年"); ax.plot(x,b,color=BLUE,lw=2,marker="o",ms=6,label="2022年")
    ax.set_xlim(-.5,5.5); ax.set_ylim(0,3200); ax.set_xticks(x,months); ax.set_yticks(np.arange(0,3001,500)); ax.tick_params(length=0,colors=WHITE,pad=7)
    for t in list(ax.get_xticklabels())+list(ax.get_yticklabels()): t.set_fontproperties(fp); t.set_fontsize(9)
    ax.set_axisbelow(True); ax.grid(axis="y",color="white",alpha=.16,linestyle=(0,(6,4)),lw=.8)
    for s in ax.spines.values(): s.set_visible(False)
    for xi,va,vb in zip(x,a,b):
        ya=va+110 if va>=vb else va-160; yb=vb-160 if vb<=va else vb+110
        ax.text(xi,ya,f"{int(va)}",ha="center",color=WHITE,fontsize=9,fontproperties=fp)
        ax.text(xi,yb,f"{int(vb)}",ha="center",color=WHITE,fontsize=9,fontproperties=fp)
    lg=ax.legend(loc="upper right",bbox_to_anchor=(.98,1.14),ncol=2,frameon=False,prop=fp)
    for t in lg.get_texts(): t.set_color(WHITE); t.set_fontsize(9)
    fig.text(.06,.92,"2022年上半年各月同比去年销量",color=WHITE,fontsize=20,fontweight="bold",fontproperties=fp,va="top"); fig.text(.06,.82,"上半年同比去年增长明显，5月份同比增长最多，增长近40%",color=WHITE,fontsize=14,fontproperties=fp,va="top"); fig.text(.045,.045,"*注：数据来源于公司销售系统，统计日期截至2022.06.30",color="#D9D9D9",fontsize=8,fontproperties=fp)
    out=base_dir()/OUTPUT; fig.savefig(out,dpi=300,facecolor=BG); print(f"图片已保存：{out.resolve()}"); plt.show()
if __name__=="__main__": main()
